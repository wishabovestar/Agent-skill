#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ppt_tool — PPT 撰写与生成工具（本机原生，中文字体验证过）

为什么要有它
--------------------------------------------------------------------------
Hermes 自带的 `powerpoint` 技能文档写了三条主线，但在本机【跑不通】：
  · 读文本   → 依赖 markitdown[pptx]        （未装）
  · 从零创建 → 依赖 pptxgenjs (npm)         （未装）
  · 视觉 QA  → 依赖 LibreOffice + pdftoppm  （未装，且本机没有 LibreOffice）
本机实际有的：python-pptx 1.0.2 / Pillow / 【PowerPoint COM（Office 12.0）】。
⇒ 本工具用「python-pptx 生成 + PowerPoint COM 渲染」，绕开全部缺失依赖。

与 `md2docx_cn.py` 对偶：那个是 MD→Word，这个是 MD→PPT，共用同一套中文排版约定。

⭐ 核心难点（与 md2docx_cn 的 `w:eastAsia` 是同一个坑的变体）
   python-pptx 里 `run.font.name = "微软雅黑"` 只写 `<a:latin>`，【中文不走 latin】，
   中文会回退到主题默认字体。要真正生效必须另写 `<a:ea typeface="微软雅黑"/>`，
   且 `a:ea` 在 CT_TextCharacterProperties 里必须【排在 a:latin 之后、a:sym 之前】。
   本文件用 insert_element_before 保证顺序，并可用 `check-fonts` 回读 XML 自证。

用法
--------------------------------------------------------------------------
    python ppt_tool.py check                          # 依赖自检
    python ppt_tool.py build in.md -o out.pptx -t navy # MD -> PPTX
    python ppt_tool.py read out.pptx                  # 提取文本 + 备注
    python ppt_tool.py render out.pptx -o png/        # PPTX -> 逐页 PNG（COM）
    python ppt_tool.py qa out.pptx                    # 内容体检
    python ppt_tool.py check-fonts out.pptx           # 回读 XML 证明 a:ea 真的写进去了

Markdown 约定（写 PPT 的源）
--------------------------------------------------------------------------
    # 一级标题        ⇒ 新的一页（第 1 个 `#` 作封面）
    ## 二级标题       ⇒ 页内小节标题
    - / * 要点        ⇒ 项目符号
    > 备注           ⇒ 演讲者备注（不进正文）
    ---              ⇒ 强制分页
    1. 2. 3.         ⇒ 有序列表
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

# ── 主题（复用 cn-application-diagrams 的既定配色，不另立标准） ────────────────
THEMES = {
    # 深蓝金正式系
    "navy": {
        "bg": "FFFFFF", "title_bg": "16365C", "accent": "C98A1B",
        "text": "2B3A4A", "muted": "5A6B7D", "card": "F2F7FC",
        "band": "1F4E79", "rule": "DBE9F7",
    },
    # 申报书黑白风格
    "gray": {
        "bg": "FFFFFF", "title_bg": "1A1A1A", "accent": "444444",
        "text": "1A1A1A", "muted": "8A8A8A", "card": "F2F2F2",
        "band": "D0D0D0", "rule": "D9D9D9",
    },
    # 深色全文
    "dark": {
        "bg": "21295C", "title_bg": "065A82", "accent": "00A896",
        "text": "F2F2F2", "muted": "CADCFC", "card": "1C7293",
        "band": "065A82", "rule": "1C7293",
    },
}

EA_FONT = "微软雅黑"      # 中文 eastAsia
LATIN_FONT = "Arial"     # 西文 latin
CJK_TITLE = "微软雅黑"

SLIDE_W = Inches(13.333)   # 16:9
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.7)


# ── ⭐ 中文字体：必须同时写 a:latin / a:ea / a:cs，且满足 schema 顺序 ──────────
# CT_TextCharacterProperties 的元素顺序：
#   a:ln, (fill), a:effectLst, a:highlight, (uLn*), a:latin, a:ea, a:cs, a:sym,
#   a:hlinkClick, a:hlinkMouseOver, a:rtl, a:extLst
_AFTER_LATIN = ("a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")
_AFTER_EA = ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def _ensure_typeface(rPr, tag: str, face: str, successors: tuple) -> None:
    el = rPr.find(qn(tag))
    if el is None:
        el = rPr.makeelement(qn(tag), {})
        rPr.insert_element_before(el, *successors)
    el.set("typeface", face)


def style_run(run, size=None, bold=None, color=None, ea=EA_FONT, latin=LATIN_FONT):
    """给 run 设字号/加粗/颜色 + 【真正生效的】中文字体。

    只设 font.name 是不够的 —— 它只写 a:latin，中文会回退到主题字体。
    """
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)
    f.name = latin                      # python-pptx 会写好 a:latin
    rPr = run._r.get_or_add_rPr()       # ⚠️ 不可用 run._r.rPr（可能为 None）
    _ensure_typeface(rPr, "a:ea", ea, _AFTER_LATIN)
    _ensure_typeface(rPr, "a:cs", latin, _AFTER_EA)
    return run


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _emit_runs(p, text, size, bold, color, ea, latin=LATIN_FONT):
    """把一行文本写进段落；行内 `**…**` 转成真加粗 run，而不是把星号打出来。

    2026-10-17 视觉 QA 实测：旧版把 `**` 原样上屏（6 页共 30 处），故加此解析。
    无标记时退化为单 run（与旧行为逐字一致）。
    """
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run()
            style_run(r, size=size, bold=bold, color=color, ea=ea, latin=latin)
            r.text = text[pos:m.start()]
        r = p.add_run()
        style_run(r, size=size, bold=True, color=color, ea=ea, latin=latin)
        r.text = m.group(1)
        pos = m.end()
    r = p.add_run()
    style_run(r, size=size, bold=bold, color=color, ea=ea, latin=latin)
    r.text = text[pos:]


def add_text(slide, x, y, w, h, text, size=16, bold=False, color="000000",
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.25,
             ea=EA_FONT, wrap=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        _emit_runs(p, ln, size, bold, color, ea)
    return box


def add_bullets(slide, x, y, w, h, items, size=16, color="2B3A4A",
                bullet_color=None, gap=10, ea=EA_FONT, numbered=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        p.space_after = Pt(gap)
        marker = f"{i+1}. " if numbered else "▪ "
        _emit_runs(p, marker + it, size, False, color, ea)
    return box


def blank(prs, theme):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor.from_string(theme["bg"])
    return s


def bar(slide, theme, x, y, w, h, color, rounded=False):
    """色块。【默认直角】—— 全幅底/标题栏用圆角会在四角露出白底（视觉 QA 实测）。"""
    from pptx.enum.shapes import MSO_SHAPE
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


# ── Markdown 解析 ───────────────────────────────────────────────────────────
def parse_md(text: str):
    """→ [{'title':..,'sub':None,'items':[..],'notes':[..],'numbered':bool}]"""
    slides, cur = [], None
    for raw in text.splitlines():
        line = raw.rstrip().replace("<br>", "\n").replace("<br/>", "\n")
        if not line.strip():
            continue
        if line.strip() == "---":
            if cur:
                slides.append(cur); cur = None
            continue
        if line.startswith("# "):
            if cur:
                slides.append(cur)
            cur = {"title": line[2:].strip(), "sub": None, "items": [], "notes": [], "numbered": False}
        elif line.startswith("## "):
            if cur is None:
                cur = {"title": "", "sub": None, "items": [], "notes": [], "numbered": False}
            cur.setdefault("subs", []).append(line[3:].strip())
            cur["items"].append(("SUB", line[3:].strip()))
        elif line.startswith("> "):
            if cur is None:
                cur = {"title": "", "sub": None, "items": [], "notes": [], "numbered": False}
            cur["notes"].append(line[2:].strip())
        elif re.match(r"^\s*[-*]\s+", line):
            if cur is None:
                continue
            cur["items"].append(("LI", re.sub(r"^\s*[-*]\s+", "", line).strip()))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            if cur is None:
                continue
            cur["numbered"] = True
            cur["items"].append(("LI", re.sub(r"^\s*\d+[.)]\s+", "", line).strip()))
        else:
            if cur is not None:
                cur["items"].append(("P", line.strip()))
    if cur:
        slides.append(cur)
    return slides


# ── build ──────────────────────────────────────────────────────────────────
def cmd_build(args) -> int:
    src = Path(args.input)
    if not src.is_file():
        print(f"❌ 找不到 {src}"); return 2
    theme = THEMES.get(args.theme)
    if theme is None:
        print(f"❌ 未知主题 {args.theme}；可选 {list(THEMES)}"); return 2
    slides = parse_md(src.read_text(encoding="utf-8", errors="replace"))
    if not slides:
        print("❌ 解析出 0 页；检查 Markdown 是否用了 `# ` 一级标题"); return 2

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    for idx, sp in enumerate(slides):
        s = blank(prs, theme)
        is_cover = idx == 0
        if is_cover:
            bar(s, theme, Emu(0), Emu(0), SLIDE_W, SLIDE_H, theme["title_bg"])
            # 文字块【垂直居中】：标题按行数估高，避免"整体偏上"（视觉 QA 实测缺陷3）
            tlines = sp["title"].count("\n") + 1
            # 封面副标题：deck 里 `## ……` 解析成 SUB（旧版只认 P ⇒ 封面子标题整行丢失，
            # 2026-10-17 视觉 QA 实测封面只有两行标题、日期不见）
            sub_items = [t for k, t in sp["items"] if k == "SUB"]
            has_sub = bool(sub_items)
            blk_h = int(Inches(0.95) * tlines + (Inches(1.05) if has_sub else 0))
            top = int((SLIDE_H - blk_h) / 2)
            add_text(s, MARGIN, Emu(top), SLIDE_W - 2 * MARGIN, Inches(0.95) * tlines,
                     sp["title"], size=40, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER,
                     line_spacing=1.35)
            if has_sub:
                sub = " · ".join(sub_items)[:80]
                if sub:
                    add_text(s, MARGIN, Emu(top + int(Inches(0.95) * tlines) + int(Inches(0.3))),
                             SLIDE_W - 2 * MARGIN, Inches(0.8),
                             sub, size=18, color=theme["accent"], align=PP_ALIGN.CENTER)
        else:
            bar(s, theme, Emu(0), Emu(0), SLIDE_W, Inches(1.15), theme["title_bg"])
            add_text(s, MARGIN, Inches(0.32), SLIDE_W - 2 * MARGIN, Inches(0.6),
                     sp["title"], size=28, bold=True, color="FFFFFF")
            li = [t for k, t in sp["items"] if k == "LI"]
            subs = [t for k, t in sp["items"] if k == "SUB"]
            paras = [t for k, t in sp["items"] if k == "P"]
            body_top = Inches(1.55)
            body_h = SLIDE_H - body_top - Inches(0.55)
            y = body_top
            if subs:
                add_text(s, MARGIN, y, SLIDE_W - 2 * MARGIN, Inches(0.5),
                         " ｜ ".join(subs), size=18, bold=True, color=theme["band"])
                y = y + Inches(0.7)
            items = li if li else paras
            if items:
                # 字号按 16:9 (13.33") 幅面定 —— 技能的 14-16pt 建议是按 4:3 写的，直接套会偏小
                # （视觉 QA 连续两轮报"字号偏小"，这里提到 24-26pt）
                n = len(items)
                size = 26 if n <= 2 else (24 if n <= 3 else (20 if n <= 5 else (17 if n <= 8 else 14)))
                est_in = n * (size * 1.32 / 72 + 10 / 72)      # 行高 + 段后（英寸）
                est = int(Inches(est_in))
                used = y - body_top
                if est < body_h - used:
                    # 0.42（而非 0.5）：只在【标题栏以下区域】居中会把块压到整页光学中心之下
                    # （视觉 QA 实测：4 条内容时观感"偏下"）
                    y = y + int((body_h - used - est) * 0.42)
                add_bullets(s, MARGIN, y, SLIDE_W - 2 * MARGIN, body_h,
                            items, size=size, color=theme["text"], numbered=sp["numbered"])
        if sp["notes"]:
            s.notes_slide.notes_text_frame.text = "\n".join(sp["notes"])

    out = Path(args.output or (src.stem + ".pptx"))
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"✅ 生成 {out}  ({out.stat().st_size:,} bytes, {len(slides)} 页, 主题 {args.theme})")
    return 0


# ── read ───────────────────────────────────────────────────────────────────
def cmd_read(args) -> int:
    prs = Presentation(args.pptx)
    n = 0
    for i, s in enumerate(prs.slides, 1):
        n += 1
        print(f"\n── 第 {i} 页 ──")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                for line in sh.text_frame.text.splitlines():
                    if line.strip():
                        print(f"   {line}")
        try:
            nt = s.notes_slide.notes_text_frame.text.strip()
            if nt:
                print(f"   [备注] {nt[:200]}")
        except Exception:
            pass
    print(f"\n共 {n} 页")
    return 0


# ── render（PowerPoint COM，绕开 LibreOffice） ───────────────────────────────
def cmd_render(args) -> int:
    try:
        import win32com.client as win32
    except Exception as e:
        print(f"❌ 需要 pywin32：{e}"); return 3
    src = Path(args.pptx).resolve()
    if not src.is_file():
        print(f"❌ 找不到 {src}"); return 2
    outdir = Path(args.output or (src.parent / (src.stem + "_png"))).resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    app = win32.Dispatch("PowerPoint.Application")
    pres = None
    try:
        pres = app.Presentations.Open(str(src), WithWindow=False)
        # 逐页导出 PNG（比整本 Export 更好控制命名）
        for i, sl in enumerate(pres.Slides, 1):
            sl.Export(str(outdir / f"slide-{i:02d}.png"), "PNG", args.width, args.height)
        made = sorted(outdir.glob("slide-*.png"))
        print(f"✅ 导出 {len(made)} 张 PNG → {outdir}")
        for f in made:
            print(f"   {f.name}  {f.stat().st_size:,} B")
        if args.pdf:
            pres.SaveAs(str(outdir / (src.stem + ".pdf")), 32)   # 32 = ppSaveAsPDF
            print(f"✅ 另存 PDF → {outdir / (src.stem + '.pdf')}")
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:
            pass
        try:
            app.Quit()
        except Exception:
            pass
    return 0


# ── qa ─────────────────────────────────────────────────────────────────────
_PLACEHOLDER = re.compile(r"(xxxx+|lorem ipsum|单击此处|此处添加|click to (add|edit))", re.I)


def cmd_qa(args) -> int:
    prs = Presentation(args.pptx)
    issues = []
    min_size = args.min_size
    for i, s in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        joined = "\n".join(texts)
        if not joined.strip():
            issues.append(f"第 {i} 页：完全空白（无任何文本）")
        m = _PLACEHOLDER.search(joined)
        if m:
            issues.append(f"第 {i} 页：残留占位符 «{m.group(0)}»")
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    sz = r.font.size.pt if r.font.size else None
                    if sz is not None and sz < min_size:
                        issues.append(f"第 {i} 页：字号 {sz}pt < {min_size}pt  «{r.text[:24]}»")
            # CJK 字体是否真的写了 a:ea
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    has_cjk = bool(re.search(r"[\u4e00-\u9fff]", r.text))
                    rPr = r._r.find(qn("a:rPr"))
                    ea = rPr.find(qn("a:ea")) if rPr is not None else None
                    if has_cjk and ea is None:
                        issues.append(f"第 {i} 页：中文 run 缺 a:ea（会回退字体） «{r.text[:20]}»")
    n = len(prs.slides._sldIdLst)
    print(f"== PPT QA == {args.pptx}")
    print(f"  页数 {n} ｜ 判据: 空页 / 占位符 / 字号<{min_size}pt / 中文缺 a:ea")
    if issues:
        print(f"  ❌ {len(issues)} 项问题：")
        for i in issues[:40]:
            print(f"     - {i}")
        return 1
    print("  ✅ 无问题")
    return 0


def cmd_check_fonts(args) -> int:
    """回读 XML 自证 a:ea 真的写进了 rPr（不靠'我设过了'）。"""
    prs = Presentation(args.pptx)
    tot = cjk = with_ea = 0
    sample = None
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    tot += 1
                    if re.search(r"[\u4e00-\u9fff]", r.text):
                        cjk += 1
                        rPr = r._r.find(qn("a:rPr"))
                        ea = rPr.find(qn("a:ea")) if rPr is not None else None
                        if ea is not None:
                            with_ea += 1
                            if sample is None:
                                sample = (r.text[:18], ea.get("typeface"))
    print(f"  run 总数 {tot} ｜ 含中文 {cjk} ｜ 其中写了 a:ea 的 {with_ea}")
    if sample:
        print(f"  样本: «{sample[0]}» → a:ea typeface = {sample[1]}")
    ok = cjk > 0 and with_ea == cjk
    print("  " + ("✅ 全部中文 run 都带 a:ea" if ok else "❌ 有中文 run 缺 a:ea（字体必回退）"))
    return 0 if ok else 1


def cmd_check(_a) -> int:
    print("== ppt_tool 依赖自检 ==")
    ok = True
    try:
        import pptx
        print(f"  ✅ python-pptx {pptx.__version__}")
    except Exception as e:
        print(f"  ❌ python-pptx {e}"); ok = False
    for mod, why in (("PIL", "缩略图"), ("win32com.client", "PowerPoint COM 渲染")):
        try:
            __import__(mod)
            print(f"  ✅ {mod}  ({why})")
        except Exception as e:
            print(f"  ⚠️ {mod} 不可用 ({why}): {type(e).__name__}")
    try:
        import win32com.client as w
        app = w.Dispatch("PowerPoint.Application"); ver = app.Version; app.Quit()
        print(f"  ✅ PowerPoint COM 可用（版本 {ver}）⇒ 渲染/PDF 不需要 LibreOffice")
    except Exception as e:
        print(f"  ❌ PowerPoint COM 不可用: {e}"); ok = False
    print("  " + ("✅ 自检通过" if ok else "❌ 有硬缺失"))
    return 0 if ok else 1


def main() -> int:
    ap = argparse.ArgumentParser(prog="ppt_tool", description="PPT 撰写与生成工具")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sub.add_parser("check", help="依赖自检").set_defaults(fn=cmd_check)

    b = sub.add_parser("build", help="Markdown -> PPTX")
    b.add_argument("input"); b.add_argument("-o", "--output")
    b.add_argument("-t", "--theme", default="navy", choices=list(THEMES))
    b.set_defaults(fn=cmd_build)

    r = sub.add_parser("read", help="提取文本+备注"); r.add_argument("pptx"); r.set_defaults(fn=cmd_read)

    d = sub.add_parser("render", help="PPTX -> 逐页 PNG（PowerPoint COM）")
    d.add_argument("pptx"); d.add_argument("-o", "--output")
    d.add_argument("--width", type=int, default=1600); d.add_argument("--height", type=int, default=900)
    d.add_argument("--pdf", action="store_true", help="同时另存 PDF")
    d.set_defaults(fn=cmd_render)

    q = sub.add_parser("qa", help="内容体检")
    q.add_argument("pptx"); q.add_argument("--min-size", type=float, default=11.0)
    q.set_defaults(fn=cmd_qa)

    c = sub.add_parser("check-fonts", help="回读 XML 证明 a:ea 写入")
    c.add_argument("pptx"); c.set_defaults(fn=cmd_check_fonts)

    a = ap.parse_args()
    return a.fn(a)


if __name__ == "__main__":
    sys.exit(main())
